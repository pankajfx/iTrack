"""
seed_users.py
─────────────
Reads user data from the Excel file (co-located with this script):
    scripts/SDWAN Installation Tracker Master User Data.xlsx

Source of truth is the per-role sheets - FE(235), FEG(23), FS(5), FSG(1),
NS(5), NSG(1) - NOT the consolidated All(270) sheet, whose Password column
holds a single placeholder value for every user.

Two modes:

  full (default)      Makes the users collection match the workbook: every sheet
                      row is upserted BY USERNAME (existing accounts keep their
                      _id), and accounts absent from the sheet are removed.
                      Passwords are hashed with Werkzeug and FEG hierarchy
                      (field_support) is derived from region → FS.

                      It used to delete every user and insert fresh documents.
                      That gave every account a new _id, and because trackers
                      store their owners by id (fe.id, noc_assignee), one reseed
                      orphaned every tracker in the database. If that has
                      happened, scripts/relink_orphans.py repairs it.

  --passwords-only    Non-destructive. Re-hashes each workbook password onto the
                      matching existing user (matched by username) and touches
                      nothing else — no deletes, no inserts, no other fields.
                      Use this to recover from a forgotten/!changed password
                      without losing accounts or hierarchy edits.

Add --dry-run to either mode to see what would happen without writing.

Usage (run from project root):
    python scripts/seed_users.py
    python scripts/seed_users.py --passwords-only
    python scripts/seed_users.py --passwords-only --dry-run

Requirements:
    pip install openpyxl werkzeug pymongo
"""

import argparse
import os
import re
import shutil
import sys
import tempfile
import zipfile
from collections import Counter
from datetime import datetime

from bson import ObjectId
from pymongo import MongoClient, UpdateOne
from werkzeug.security import check_password_hash, generate_password_hash
from dotenv import load_dotenv

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# Load environment variables from the project-root .env (parent of scripts/).
load_dotenv(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '.env'))

# ── Config ────────────────────────────────────────────────────────────────────
MONGO_URI = os.environ.get('MONGO_URI', 'mongodb://localhost:27017/sdwan_tracker')
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(SCRIPT_DIR, 'SDWAN Installation Tracker Master User Data.xlsx')

# ── Role mapping: Excel label → app constant ──────────────────────────────────
ROLE_MAP = {
    'Field Support Group':  'FIELD_SUPPORT_GROUP',
    'Field Support':        'FIELD_SUPPORT',
    'Field Engineer Group': 'FIELD_ENGINEER_GROUP',
    'Field Engineer':       'FIELD_ENGINEER',
    'NOC Support Group':    'NOC_SUPPORT_GROUP',
    'NOC Support':          'NOC_SUPPORT',
    'Analytics':            'ANALYTICS',
}

# ── Read Excel ────────────────────────────────────────────────────────────────
# Excel writes autofilter values that openpyxl's validator rejects (this workbook
# has <customFilter val=" ">), which makes load_workbook() raise before a single row
# is read. The filters are irrelevant to us, so strip them from a throwaway copy.
_AUTOFILTER_RE = re.compile(
    rb'<autoFilter\b[^>]*/>|<autoFilter\b[^>]*>.*?</autoFilter>', re.S)


def _workbook_without_filters(path):
    """Copy the .xlsx to a temp file with every <autoFilter> removed."""
    tmp_dir = tempfile.mkdtemp(prefix='seed_users_')
    tmp = os.path.join(tmp_dir, os.path.basename(path))
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml'):
                data = _AUTOFILTER_RE.sub(b'', data)
            zout.writestr(item, data)
    return tmp_dir, tmp


# The consolidated "All(270)" sheet carries PLACEHOLDER passwords - a single value
# repeated for every user. The per-role sheets (FE(235), FEG(23), FS(5), FSG(1),
# NS(5), NSG(1)) hold the real per-user passwords and agree with the consolidated
# sheet on every other column, so they are the source of truth. Seeding from
# All(270) is what put "test123" on all 270 accounts and broke every real login.
ROLE_SHEET_RE = re.compile('^(FE|FEG|FS|FSG|NS|NSG)[(][0-9]+[)]$')


def _sheet_rows(ws):
    """Rows of one sheet as dicts keyed by its header row; blank rows dropped."""
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = list(next(rows_iter))
    except StopIteration:
        return []
    out = []
    for values in rows_iter:
        row = dict(zip(headers, values))
        if row.get('Role'):          # skip blank trailing rows
            out.append(row)
    return out


def _warn_if_placeholder_passwords(rows):
    """Flag a sheet whose rows all share one password - the All(270) failure mode."""
    passwords = {pw for pw in (row_credentials(r)[1] for r in rows) if pw}
    if len(rows) > 10 and len(passwords) == 1:
        print("")
        print("  *** WARNING: all %d rows share one password." % len(rows))
        print("      That is a placeholder sheet, not real credentials.")
        print("      Seeding this would lock every user out of their real password.")


def read_excel_rows():
    """Return user rows, preferring the per-role sheets over the consolidated one."""
    tmp_dir, tmp_path = _workbook_without_filters(EXCEL_PATH)
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        role_sheets = [n for n in wb.sheetnames if ROLE_SHEET_RE.match(n)]
        if role_sheets:
            rows, seen = [], {}
            for name in role_sheets:
                for row in _sheet_rows(wb[name]):
                    username = str(row.get('Username', '')).strip()
                    if username in seen:
                        print("  WARNING: duplicate username %r in %s (kept %s)"
                              % (username, name, seen[username]))
                        continue
                    seen[username] = name
                    rows.append(row)
            print("  Source sheets: %s" % ', '.join(role_sheets))
        else:
            name = 'All(270)' if 'All(270)' in wb.sheetnames else wb.sheetnames[0]
            rows = _sheet_rows(wb[name])
            print("  Source sheet: %s (no per-role sheets found)" % name)
        wb.close()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    _warn_if_placeholder_passwords(rows)
    return rows


def row_credentials(row):
    """(username, plaintext password) for a sheet row; password is None if blank.

    A blank password used to default to the username - a credential anyone
    who can read the login page's name list could guess. Rows without a
    password are now skipped and reported instead.
    """
    username = str(row.get('Username', '')).strip()
    raw = row.get('Password')
    password = str(raw).strip() if raw is not None else ''
    return username, (password or None)


# ── Build MongoDB documents ───────────────────────────────────────────────────
def build_user_docs(rows):
    now = datetime.utcnow()

    # Build region → FS username lookup (used to set field_support on FEG docs)
    region_to_fs = {}
    for row in rows:
        if row.get('Role') == 'Field Support':
            region = row.get('Region')
            if region:
                region_to_fs[region] = str(row['Username']).strip()

    docs = []
    skipped = []

    for row in rows:
        excel_role = str(row.get('Role', '')).strip()
        role = ROLE_MAP.get(excel_role)
        if not role:
            skipped.append(f"Unknown role '{excel_role}' — username: {row.get('Username')}")
            continue

        username, password_plain = row_credentials(row)
        if not password_plain:
            skipped.append(f"No password in the sheet - username: {username}")
            continue
        name           = str(row.get('Name', username)).strip()
        zone           = str(row['Zone']).strip() if row.get('Zone') else 'India'
        region         = str(row['Region']).strip() if row.get('Region') else None
        group          = str(row['Group']).strip() if row.get('Group') else None   # FE → FEG name
        state          = str(row['State']).strip() if row.get('State') else None   # FE → location
        email          = str(row['Email']).strip() if row.get('Email') else None
        # Contact may be stored as float (e.g. 6302828144.0) — normalise to string
        raw_contact    = row.get('Contact')
        contact        = str(int(raw_contact)) if isinstance(raw_contact, float) else (
                         str(raw_contact).strip() if raw_contact else None)

        doc = {
            'username':   username,
            'name':       name,
            'password':   generate_password_hash(password_plain),
            'role':       role,
            'active':     True,
            'created_at': now,
            'updated_at': now,
            'zone':       zone,
        }

        if region:
            doc['region'] = region

        # Role-specific hierarchy / contact fields
        if role == 'FIELD_ENGINEER':
            if group:
                doc['field_engineer_group'] = group
            if state:
                doc['location'] = state
            if email:
                doc['email'] = email
            if contact:
                doc['contact'] = contact

        elif role == 'FIELD_ENGINEER_GROUP':
            # Derive parent FS from region
            if region and region in region_to_fs:
                doc['field_support'] = region_to_fs[region]

        docs.append(doc)

    if skipped:
        print(f"\n  WARNINGS ({len(skipped)}):")
        for msg in skipped:
            print(f"    {msg}")

    return docs


# ── Password-only reset (non-destructive) ─────────────────────────────────────
def reset_passwords(db, rows, dry_run=False):
    """Re-hash each workbook password onto the matching existing user.

    Only the `password` field (and `updated_at`) is touched — accounts, roles and
    hierarchy stay exactly as they are. Users absent from the workbook are left
    alone; workbook rows with no matching username are reported, not created.
    """
    now = datetime.utcnow()
    updated = unchanged = not_found = 0
    missing = []

    for row in rows:
        username, password_plain = row_credentials(row)
        if not username or not password_plain:
            continue

        user = db.users.find_one({'username': username}, {'_id': 1, 'password': 1})
        if not user:
            not_found += 1
            missing.append(username)
            continue

        # Skip the write when the stored hash already accepts this password —
        # scrypt hashes are salted, so comparing hashes directly would never match.
        if user.get('password') and check_password_hash(user['password'], password_plain):
            unchanged += 1
            continue

        if not dry_run:
            db.users.update_one(
                {'_id': user['_id']},
                {'$set': {'password': generate_password_hash(password_plain),
                          'updated_at': now}}
            )
        updated += 1

    verb = "would be reset" if dry_run else "reset"
    print("")
    print(f"  Passwords {verb}      : {updated}")
    print(f"  Already correct        : {unchanged}")
    print(f"  In sheet, not in DB    : {not_found}")
    if missing:
        preview = ', '.join(missing[:10])
        print(f"    {preview}{' …' if len(missing) > 10 else ''}")

    total = db.users.count_documents({})
    print(f"  Users in DB (untouched count): {total}")


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description='Seed users from the master workbook. See the module docstring '
                    'for full details.')
    parser.add_argument('--passwords-only', action='store_true',
                        help="Only re-hash passwords onto existing users. "
                             "No deletes, no inserts, no other fields changed.")
    parser.add_argument('--dry-run', action='store_true',
                        help='Report what would change without writing anything.')
    args = parser.parse_args()

    print(f"Excel source: {EXCEL_PATH}")
    if not os.path.exists(EXCEL_PATH):
        print("ERROR: Excel file not found.")
        sys.exit(1)

    rows = read_excel_rows()
    print(f"  Read {len(rows)} rows from Excel")

    client = MongoClient(MONGO_URI)
    db     = client.get_default_database()

    if args.passwords_only:
        print("")
        print(f"Mode: PASSWORDS ONLY{' (dry run)' if args.dry_run else ''} - "
              f"no user will be deleted or created.")
        reset_passwords(db, rows, dry_run=args.dry_run)
        client.close()
        return

    docs = build_user_docs(rows)
    print(f"  Built {len(docs)} user documents")

    sheet_usernames = [d['username'] for d in docs]
    if args.dry_run:
        existing = db.users.count_documents({'username': {'$in': sheet_usernames}})
        prune = db.users.count_documents({'username': {'$nin': sheet_usernames}})
        print("")
        print(f"DRY RUN - would update {existing} existing user(s) in place, "
              f"create {len(docs) - existing}, and remove {prune} not in the sheet.")
        client.close()
        return

    # Ensure collection + unique index on username
    if 'users' not in db.list_collection_names():
        db.create_collection('users')
        print("  Created 'users' collection")
    # Same name create_indexes.py uses - an unnamed create_index() here would
    # take the default name 'username_1' and then collide with it (IndexOptionsConflict).
    db.users.create_index('username', unique=True, name='users_username_unique')

    # Upsert by username so every existing account keeps its _id - see the
    # module docstring for why delete-and-insert was destructive to trackers.
    print(f"\nUpserting {len(docs)} users by username...")
    ops = []
    for d in docs:
        d = dict(d)
        created_at = d.pop('created_at')
        ops.append(UpdateOne({'username': d['username']},
                             {'$set': d, '$setOnInsert': {'created_at': created_at}},
                             upsert=True))
    result = db.users.bulk_write(ops, ordered=False)
    print(f"  Updated {result.modified_count}, unchanged "
          f"{result.matched_count - result.modified_count}, created {result.upserted_count}")

    gone = [str(u['_id']) for u in db.users.find({'username': {'$nin': sheet_usernames}}, {'_id': 1})]
    if gone:
        still_owned = db.trackers.count_documents(
            {'$or': [{'fe.id': {'$in': gone}}, {'noc_assignee': {'$in': gone}}]})
        removed = db.users.delete_many({'_id': {'$in': [ObjectId(g) for g in gone]}}).deleted_count
        print(f"  Removed {removed} user(s) not in the sheet")
        if still_owned:
            print(f"  NOTE: {still_owned} tracker(s) belonged to removed users and now "
                  f"need reassigning - run scripts/relink_orphans.py to list them.")

    # Summary
    counts = Counter(d['role'] for d in docs)
    print("\nUsers by role:")
    for role, count in sorted(counts.items()):
        print(f"  {role:<30} {count}")
    print(f"\nTotal: {len(docs)} users seeded successfully.")

    client.close()


if __name__ == '__main__':
    main()
